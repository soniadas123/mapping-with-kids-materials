"""Build the Exercise 1 intro deck: Intro_to_Mapping.pptx

Run from anywhere inside the mapping_kids conda env:
    python build_intro_pptx.py

Photos in images/ were cropped from the 2021 "Mapping for Kids" blog slides.
Every diagram is drawn with plain PowerPoint shapes, so it can be edited later.
"""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
IMAGES = HERE / "images"
OUTPUT = HERE / "Intro_to_Mapping.pptx"

# Colours: deep space navy, earth green, sky blue, sun amber
NAVY = RGBColor(0x0F, 0x23, 0x40)
GREEN = RGBColor(0x2E, 0x7D, 0x4F)
BLUE = RGBColor(0x1F, 0x6F, 0xB2)
AMBER = RGBColor(0xE8, 0xA3, 0x17)
RED = RGBColor(0xD9, 0x48, 0x2B)
TEXT = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER_TINT = RGBColor(0xFE, 0xF3, 0xC7)
GREEN_TINT = RGBColor(0xE6, 0xF4, 0xEA)
BLUE_TINT = RGBColor(0xE7, 0xF0, 0xFA)
GREY_TINT = RGBColor(0xF3, 0xF4, 0xF6)

FONT = "Calibri"


# ---------- small helpers ----------

def new_slide(prs, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY if dark else WHITE
    return slide


def add_text(slide, x, y, w, h, lines, size=22, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullets=False):
    """Add a text box. `lines` is a string or a list of strings (one paragraph each)."""
    if isinstance(lines, str):
        lines = [lines]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(10)
        if bullets:
            make_bullet(p)
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def make_bullet(paragraph):
    """Turn a paragraph into a round bullet with a hanging indent."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.35)))
    pPr.set("indent", str(-Inches(0.3)))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))


def add_title(slide, text, color=NAVY):
    add_text(slide, 0.6, 0.4, 12.1, 0.9, text, size=36, bold=True, color=color)


def add_shape(slide, kind, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    shape.shadow.inherit = False
    return shape


def add_card(slide, x, y, w, h, label, body, tint, label_color, size=18):
    """A rounded, tinted box with a small coloured label and body text."""
    card = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, tint)
    card.adjustments[0] = 0.08
    add_text(slide, x + 0.25, y + 0.2, w - 0.5, 0.4, label.upper(), size=14,
             bold=True, color=label_color)
    add_text(slide, x + 0.25, y + 0.6, w - 0.5, h - 0.75, body, size=size)
    return card


def word_of_the_day(slide, x, y, w, h, word, meaning):
    add_card(slide, x, y, w, h, "Word of the day", f"{word}: {meaning}",
             AMBER_TINT, RGBColor(0x92, 0x40, 0x0E))


def fun_fact(slide, x, y, w, h, text):
    add_card(slide, x, y, w, h, "Fun fact", text, GREEN_TINT, GREEN)


def add_picture(slide, path, x, y, w=None, h=None):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
    pic.line.color.rgb = RGBColor(0xD1, 0xD5, 0xDB)
    pic.line.width = Pt(1)
    return pic


def add_caption(slide, x, y, w, text, align=PP_ALIGN.LEFT, color=MUTED):
    add_text(slide, x, y, w, 0.4, text, size=12, color=color, align=align)


def add_arrow(slide, x1, y1, x2, y2, color, width=3, dashed=False):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                      Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    if dashed:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = line.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))
    return line


def add_polyline(slide, points, color, width=4, closed=False, fill=None):
    """Draw a line (or a closed polygon) through a list of (x, y) points in inches."""
    emu = [(Inches(px), Inches(py)) for px, py in points]
    builder = slide.shapes.build_freeform(emu[0][0], emu[0][1], scale=1.0)
    builder.add_line_segments(emu[1:], close=closed)
    shape = builder.convert_to_shape()
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    shape.shadow.inherit = False
    return shape


def add_dot(slide, cx, cy, d, color, outline=WHITE):
    dot = add_shape(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, color, line=outline)
    dot.line.width = Pt(2)
    return dot


def number_dot(slide, cx, cy, text, fill=GREEN):
    """A small numbered circle, used to match labels to spots on a picture."""
    dot = add_dot(slide, cx, cy, 0.48, fill)
    tf = dot.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = WHITE
    return dot


def add_badge(slide, x, y, text, fill, size=16, w=None):
    """A small rounded label, like a sticker on a photo."""
    w = w or 0.25 + 0.12 * len(text)
    badge = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.45, fill)
    tf = badge.text_frame
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = WHITE
    return badge


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------- the slides ----------

def slide_title(prs):
    s = new_slide(prs, dark=True)
    add_text(s, 0.8, 1.5, 7.2, 0.5, "MAPPING WITH KIDS  |  BEFORE EXERCISE 1",
             size=16, bold=True, color=AMBER)
    add_text(s, 0.8, 2.1, 7.2, 2.0, "Seeing Earth from Above", size=54, bold=True,
             color=WHITE)
    add_text(s, 0.8, 4.2, 7.0, 1.2,
             "Satellites, GIS, points, lines and polygons, and how a map is made",
             size=24, color=RGBColor(0xCB, 0xD5, 0xE1))
    add_picture(s, HERE / "screenshot" / "Map-of-India.png", 8.3, 0.6, h=6.3)
    add_notes(s, "Welcome everyone. Today we find out how satellites take pictures of "
                 "the Earth, and how GIS software turns those pictures into maps. "
                 "By the end of the session you will build the map on the right yourself.")


def slide_journey(prs):
    s = new_slide(prs)
    add_title(s, "Today's journey")
    steps = [("1", "Remote sensing", BLUE), ("2", "Satellites and change", BLUE),
             ("3", "GIS and layers", GREEN), ("4", "Points, lines, polygons", GREEN),
             ("5", "Parts of a map", AMBER), ("6", "Make your own map", RED)]
    x = 0.75
    for i, (num, label, colour) in enumerate(steps):
        circle = add_shape(s, MSO_SHAPE.OVAL, x + 0.45, 2.6, 1.1, 1.1, colour)
        tf = circle.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = num
        run.font.size = Pt(32)
        run.font.bold = True
        run.font.name = FONT
        run.font.color.rgb = WHITE
        add_text(s, x, 3.9, 2.0, 1.0, label, size=18, bold=True,
                 align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_arrow(s, x + 1.65, 3.15, x + 2.25, 3.15, MUTED, width=2)
        x += 2.0
    add_text(s, 0.75, 5.5, 11.8, 0.8,
             "At the end you will open GeoLibre and build a Map of India.",
             size=22, color=MUTED, align=PP_ALIGN.CENTER)
    add_notes(s, "Walk through the six stops. Steps 1 to 5 are ideas, step 6 is hands-on "
                 "in GeoLibre. Ask: has anyone used Google Maps or Google Earth before?")


def slide_remote_sensing(prs):
    s = new_slide(prs)
    add_title(s, "What is remote sensing?")
    add_text(s, 0.6, 1.5, 7.0, 1.6,
             "Remote sensing means learning about something without touching it. "
             "Looking at a bird through binoculars is a simple kind of remote sensing.",
             size=22)
    examples = [("Weather satellite", "Watches a cyclone form over the Bay of Bengal", BLUE),
                ("Drone", "Flies over a farm to see which crops need water", GREEN),
                ("Space station camera", "Photographs a coastline as it passes overhead", AMBER)]
    y = 3.35
    for name, text, colour in examples:
        add_dot(s, 0.95, y + 0.4, 0.45, colour)
        add_text(s, 1.4, y, 6.2, 0.45, name, size=20, bold=True)
        add_text(s, 1.4, y + 0.42, 6.2, 0.5, text, size=18, color=MUTED)
        y += 1.1
    word_of_the_day(s, 8.3, 1.6, 4.4, 2.0, "Remote sensing",
                    "gathering information about a place from far away, "
                    "without touching it.")
    add_notes(s, "Ask: what else can you sense without touching? (Hearing thunder, "
                 "smelling food, seeing stars.) Satellites do the same, from hundreds of "
                 "kilometres up.")


def slide_how_satellites_see(prs):
    s = new_slide(prs)
    add_title(s, "How do satellites see the Earth?")
    # Sun, Earth and satellite, redrawn from the 2021 blog slide
    sun = add_shape(s, MSO_SHAPE.OVAL, 10.6, 1.4, 1.5, 1.5, AMBER)
    add_text(s, 10.6, 2.95, 1.5, 0.4, "Sun", size=18, bold=True, align=PP_ALIGN.CENTER)
    earth = add_shape(s, MSO_SHAPE.OVAL, 5.3, 3.3, 3.0, 3.0, BLUE)
    add_shape(s, MSO_SHAPE.OVAL, 5.9, 3.9, 1.0, 0.7, GREEN)
    add_shape(s, MSO_SHAPE.OVAL, 7.0, 5.0, 0.8, 0.6, GREEN)
    add_text(s, 5.3, 6.35, 3.0, 0.4, "Earth", size=18, bold=True, align=PP_ALIGN.CENTER)
    # satellite: body with two solar panels
    add_shape(s, MSO_SHAPE.RECTANGLE, 3.0, 1.6, 0.9, 0.24, BLUE, line=NAVY)
    add_shape(s, MSO_SHAPE.RECTANGLE, 4.25, 1.6, 0.9, 0.24, BLUE, line=NAVY)
    add_shape(s, MSO_SHAPE.RECTANGLE, 3.9, 1.47, 0.35, 0.5, NAVY)
    add_text(s, 3.0, 2.05, 2.15, 0.4, "Satellite", size=18, bold=True,
             align=PP_ALIGN.CENTER)
    add_arrow(s, 10.65, 2.5, 8.15, 4.1, AMBER, width=4)
    add_arrow(s, 5.75, 3.6, 4.35, 2.45, GREEN, width=3, dashed=True)
    add_text(s, 9.3, 3.9, 3.6, 0.8, "1. Sunlight shines on the Earth", size=18,
             color=TEXT)
    add_text(s, 0.6, 3.0, 4.5, 1.2, "2. Light bounces off the ground and the "
             "satellite's camera catches it", size=18)
    add_card(s, 0.6, 4.6, 4.3, 1.7, "Seeing more than our eyes",
             "Many satellites also see infrared light. Healthy plants shine brightly "
             "in infrared.", BLUE_TINT, BLUE, size=16)
    add_notes(s, "The satellite does not shine a torch. The Sun lights up the Earth, light "
                 "bounces back up, and the satellite camera records it, like a phone camera. "
                 "Infrared is light our eyes cannot see. Scientists use it to check if crops "
                 "and forests are healthy.")


def slide_satellites(prs):
    s = new_slide(prs)
    add_title(s, "Natural and artificial satellites")
    add_text(s, 0.6, 1.35, 12.1, 0.6,
             "A satellite is anything that travels around a planet.", size=22, color=MUTED)
    add_picture(s, IMAGES / "moon.png", 0.6, 2.2, h=3.9)
    add_text(s, 0.6, 6.2, 3.3, 0.5, "Natural: the Moon", size=20, bold=True)
    add_picture(s, IMAGES / "sputnik.png", 4.25, 2.2, h=3.9)
    add_text(s, 4.25, 6.2, 5.4, 0.5, "Artificial: Sputnik-1 (1957)", size=20, bold=True)
    fun_fact(s, 9.95, 2.2, 2.8, 2.4,
             "Sputnik-1 was the first artificial satellite. It was about the size of "
             "a beach ball.")
    add_notes(s, "Natural satellites are made by nature, like our Moon. Artificial "
                 "satellites are made by people. Sputnik-1 was launched by the Soviet Union "
                 "on 4 October 1957 and went around the Earth every 96 minutes. Today "
                 "thousands of satellites circle the Earth.")


def slide_india_in_space(prs):
    s = new_slide(prs)
    add_title(s, "India goes to space")
    add_picture(s, IMAGES / "sarabhai.png", 0.6, 1.5, h=4.4)
    add_caption(s, 0.6, 5.95, 3.4, "Dr Vikram Sarabhai, father of the Indian space programme")
    events = [("1962", "INCOSPAR is set up, led by Dr Vikram Sarabhai"),
              ("1963", "First rocket launch from Thumba, Kerala"),
              ("1969", "ISRO, the Indian Space Research Organisation, is formed"),
              ("1975", "Aryabhata, India's first satellite, goes to space"),
              ("2023", "Chandrayaan-3 lands near the Moon's south pole")]
    # a vertical timeline
    x_line = 4.85
    add_shape(s, MSO_SHAPE.RECTANGLE, x_line - 0.03, 1.75, 0.06, 4.6, RGBColor(0xD1, 0xD5, 0xDB))
    y = 1.5
    for year, text in events:
        add_dot(s, x_line, y + 0.3, 0.32, GREEN if year != "2023" else AMBER)
        add_text(s, 5.3, y, 1.1, 0.55, year, size=22, bold=True, color=GREEN)
        add_text(s, 6.45, y + 0.02, 6.3, 0.9, text, size=20)
        y += 1.0
    add_notes(s, "India's space journey started in 1962 with INCOSPAR. The first rockets "
                 "went up from Thumba near Thiruvananthapuram. ISRO was formed in 1969 "
                 "and is headquartered here in Bengaluru. Indian satellites like Cartosat "
                 "and Resourcesat take photos of India that are used to make maps.")


def slide_satellite_photo(prs):
    s = new_slide(prs)
    add_title(s, "What can you spot from space?")
    add_picture(s, IMAGES / "devanahalli_fort.png", 0.6, 1.4, h=5.5)
    add_text(s, 9.3, 1.5, 3.5, 0.5, "Devanahalli Fort, near Bengaluru", size=20,
             bold=True)
    add_text(s, 9.3, 2.3, 3.5, 3.0,
             ["The fort walls", "A highway", "Houses and streets", "Open fields"],
             size=20, bullets=True)
    add_caption(s, 9.3, 6.4, 3.5, "Satellite image: Google Earth, Maxar 2021")
    add_notes(s, "Give students a minute to find things before revealing the list. The "
                 "curved shape is the old fort wall. The long straight line on the left "
                 "is the highway to the airport. Ask: how can you tell a road from a wall?")


def slide_change(prs):
    s = new_slide(prs)
    add_title(s, "Satellites show how places change")
    add_picture(s, IMAGES / "change.png", 0.6, 1.5, h=3.7)
    add_badge(s, 0.75, 1.65, "2011", NAVY)
    add_picture(s, IMAGES / "tnh_apmc.png", 6.85, 1.5, h=3.7)
    add_badge(s, 7.0, 1.65, "2021", NAVY)
    add_text(s, 0.6, 5.4, 12.1, 0.6, "The same neighbourhood near Bengaluru, ten years apart."
             " What changed?", size=22, bold=True)
    add_text(s, 0.6, 6.05, 12.1, 0.5,
             "Try it yourself: earthengine.google.com/timelapse", size=18, color=BLUE)
    add_notes(s, "In 2011 the area was fields and an orchard. By 2021 there is a big market "
                 "yard (APMC) and an apartment complex. Satellites have photographed Earth "
                 "for over 50 years (Landsat started in 1972), so we can compare any place "
                 "over time. If there is time, open the Timelapse site and look at your "
                 "own city.")


def slide_land_cover(prs):
    s = new_slide(prs)
    add_title(s, "This is called land cover change")
    add_text(s, 0.6, 1.35, 12.1, 1.0,
             "Land cover is whatever covers the ground. When it changes over the years, "
             "we call it land cover change.", size=22)
    # the kinds of land cover a satellite image can be sorted into
    classes = [("Built-up", RGBColor(0x9C, 0xA3, 0xAF)), ("Vegetation", GREEN),
               ("Water", BLUE), ("Open land", RGBColor(0xC8, 0xA2, 0x6B))]
    y = 2.6
    for name, colour in classes:
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, y, 0.6, 0.6, colour)
        add_text(s, 1.4, y + 0.08, 3.0, 0.5, name, size=20, bold=True)
        y += 0.85
    add_text(s, 5.0, 2.45, 7.5, 0.5, "Who uses it?", size=24, bold=True, color=NAVY)
    users = [("Urban planners", "See where a city is growing, and where to keep parks and lakes", BLUE),
             ("Farmers and scientists", "Track crops, forests and droughts", GREEN),
             ("Disaster teams", "Find areas that flood easily", RED)]
    y = 3.1
    for name, text, colour in users:
        add_dot(s, 5.25, y + 0.4, 0.45, colour)
        add_text(s, 5.7, y, 7.0, 0.45, name, size=20, bold=True)
        add_text(s, 5.7, y + 0.42, 7.0, 0.5, text, size=18, color=MUTED)
        y += 1.1
    add_notes(s, "Go back to the last slide: in 2011 that area was an orchard, and by 2021 it "
                 "was buildings and a market yard. That is land cover change, from vegetation "
                 "to built-up. Computers can sort every pixel of a satellite image into classes "
                 "like these, then count how much of each there is, year after year. Urban "
                 "planners use these maps to decide where a city should grow and which lakes, "
                 "parks and forests must be protected.")


def slide_bengaluru_story(prs):
    s = new_slide(prs)
    add_title(s, "Bengaluru: from Garden City to concrete city")
    data = CategoryChartData()
    data.categories = ["Built-up", "Vegetation", "Water"]
    data.add_series("1973", (8, 68, 3.4))
    data.add_series("2013", (74, 15, 1))
    frame = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.4),
                               Inches(6.6), Inches(4.9), data)
    chart = frame.chart
    chart.font.name = FONT
    chart.font.size = Pt(16)
    chart.font.color.rgb = TEXT
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.has_title = True
    chart.chart_title.text_frame.text = "% of Bengaluru's area"
    title_font = chart.chart_title.text_frame.paragraphs[0].runs[0].font
    title_font.size = Pt(16)
    title_font.bold = False
    title_font.color.rgb = MUTED
    for series, colour in zip(chart.plots[0].series, [RGBColor(0xB8, 0xC0, 0xCC), NAVY]):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = colour
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.number_format = '0"%"'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    # Write each bar's label by hand, because IISc gives some values as "under 15%"
    labels = {"1973": ["8%", "68%", "3.4%"], "2013": ["74%", "under 15%", "under 1%"]}
    for series in plot.series:
        for point, text in zip(series.points, labels[series.name]):
            tf = point.data_label.text_frame
            tf.text = text
            tf.paragraphs[0].runs[0].font.size = Pt(14)
            point.data_label.position = XL_LABEL_POSITION.OUTSIDE_END
    value_axis = chart.value_axis
    value_axis.visible = False
    value_axis.has_major_gridlines = False
    value_axis.maximum_scale = 85
    chart.category_axis.format.line.color.rgb = RGBColor(0xD1, 0xD5, 0xDB)
    add_caption(s, 0.6, 6.4, 6.6,
                "Source: Ramachandra et al., Energy & Wetlands Research Group, IISc (2017)")
    add_text(s, 7.8, 1.5, 4.9, 0.5, "Why it matters", size=24, bold=True, color=NAVY)
    add_text(s, 7.8, 2.2, 4.9, 2.6,
             ["Fewer lakes and wetlands: rain cannot soak in, so streets flood",
              "Fewer trees: the city gets hotter",
              "Less water in the ground: borewells go dry"], size=20, bullets=True)
    fun_fact(s, 7.8, 4.85, 4.9, 1.6,
             "Scientists at IISc, right here in Bengaluru, measured this using satellite "
             "images from 1973 onwards.")
    add_notes(s, "Bengaluru was once called the Garden City. Scientists at the Indian "
                 "Institute of Science (IISc) used Landsat satellite images from 1973 to 2013 "
                 "to measure the change. Built-up area grew from about 8% of the city to about "
                 "74%. Vegetation fell from about 68% to under 15%. Water bodies fell from "
                 "3.4% to under 1%. This is exactly the remote sensing we just learned about. "
                 "Ask: what could we do to protect our lakes and trees?")


def slide_gis(prs):
    s = new_slide(prs)
    add_title(s, "What is GIS?")
    add_text(s, 0.6, 1.5, 5.8, 2.2,
             "GIS stands for Geographic Information System. It is software that "
             "stacks information about places into layers, like a sandwich.",
             size=22)
    word_of_the_day(s, 0.6, 4.0, 5.6, 1.6, "Layer",
                    "one slice of the map sandwich, such as a photo, borders or names.")
    # three stacked layers, top one first in the list
    layers = [("City names", AMBER), ("Country outline", GREEN), ("Satellite photo", BLUE)]
    y = 1.7
    for name, colour in layers:
        layer = add_shape(s, MSO_SHAPE.PARALLELOGRAM, 7.0, y, 3.6, 1.2, colour, line=WHITE)
        layer.line.width = Pt(2)
        add_text(s, 10.75, y + 0.35, 2.2, 0.5, name, size=20, bold=True)
        y += 1.45
    add_notes(s, "A GIS map is like a sandwich or a stack of see-through sheets. The bottom "
                 "layer is usually a satellite photo, called the basemap. On top we add "
                 "borders, roads, rivers and names. In Exercise 1 you will add two layers: "
                 "India's outline and the state outlines.")


def slide_points_lines_polygons(prs):
    s = new_slide(prs)
    add_title(s, "Points, lines and polygons")
    add_text(s, 0.6, 1.3, 12.1, 0.6,
             "In GIS, almost everything on a map is drawn with just three shapes.",
             size=22, color=MUTED)
    cards = [("Point", "A single spot", ["Your school", "A temple or a fort"], BLUE),
             ("Line", "A path from place to place", ["A road or railway", "A river"], GREEN),
             ("Polygon", "A closed shape with an inside", ["A park or a lake", "A state"], AMBER)]
    x = 0.6
    for name, meaning, examples, colour in cards:
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.1, 3.85, 4.7, GREY_TINT).adjustments[0] = 0.06
        if name == "Point":
            for cx, cy in [(x + 1.2, 3.0), (x + 2.2, 2.75), (x + 2.8, 3.4)]:
                add_dot(s, cx, cy, 0.35, colour)
        elif name == "Line":
            add_polyline(s, [(x + 0.5, 3.6), (x + 1.3, 2.7), (x + 2.3, 3.3), (x + 3.3, 2.6)],
                         colour, width=6)
        else:
            add_polyline(s, [(x + 0.8, 2.6), (x + 2.4, 2.5), (x + 3.1, 3.2), (x + 2.2, 3.75),
                             (x + 1.0, 3.55)], colour, width=4, closed=True,
                         fill=AMBER_TINT)
        add_text(s, x + 0.3, 4.05, 3.3, 0.55, name, size=26, bold=True, color=colour)
        add_text(s, x + 0.3, 4.6, 3.3, 0.5, meaning, size=18, color=MUTED)
        add_text(s, x + 0.3, 5.2, 3.3, 1.4, examples, size=18, bullets=True)
        x += 4.1
    add_notes(s, "Point: one location, like a pin on a map. Line: connects points, like a "
                 "road. Polygon: a line that closes back on itself, so it has an inside and "
                 "an area, like a lake or a state. Ask: is India a point, a line or a "
                 "polygon? (It depends on the scale. On a world map Bengaluru is a point, "
                 "but zoom in and it becomes a polygon.)")


def slide_spot_the_shapes(prs):
    s = new_slide(prs)
    add_title(s, "Spot the shapes")
    # The photo is 1034 x 646 pixels; place it 8.4 inches wide
    px_x, px_y, width = 0.6, 1.4, 8.4
    k = width / 1034  # inches per pixel

    def at(px, py):
        return (px_x + px * k, px_y + py * k)

    add_picture(s, IMAGES / "tnh_apmc.png", px_x, px_y, w=width)
    add_polyline(s, [at(0, 22), at(460, 92), at(1034, 190)], RED, width=5)
    add_polyline(s, [at(205, 385), at(445, 410), at(440, 585), at(210, 590)], AMBER,
                 width=5, closed=True)
    cx, cy = at(615, 357)
    add_dot(s, cx, cy, 0.32, BLUE)
    add_badge(s, *at(640, 90), "Line", RED)
    add_badge(s, *at(240, 300), "Polygon", AMBER)
    add_badge(s, cx + 0.25, cy - 0.2, "Point", BLUE)
    add_text(s, 9.4, 1.5, 3.4, 0.5, "Can you guess?", size=24, bold=True)
    add_text(s, 9.4, 2.15, 3.4, 3.0,
             ["What is the red line?", "What is inside the yellow polygon?",
              "What is the blue point?"], size=20, bullets=True)
    add_card(s, 9.4, 4.6, 3.3, 1.9, "Your turn",
             "Find one more point, line and polygon in this picture.", BLUE_TINT, BLUE,
             size=16)
    add_notes(s, "Answers: the red line follows the railway track. The yellow polygon is an "
                 "apartment complex. The blue point is a single big tree. The large grey "
                 "buildings at the top are the APMC market yard, which could be another "
                 "polygon. The road running down the middle is another line.")


def slide_raster_vector(prs):
    s = new_slide(prs)
    add_title(s, "Two kinds of map data")
    # Raster: a grid of coloured squares
    colours = {"w": BLUE, "g": GREEN, "s": RGBColor(0xC8, 0xA2, 0x6B)}
    grid = ["ssssgggg",
            "sssggggg",
            "ssggggww",
            "sgggggww",
            "ggggwwww",
            "gggwwwww"]
    cell = 0.55
    for r, row in enumerate(grid):
        for c, key in enumerate(row):
            sq = add_shape(s, MSO_SHAPE.RECTANGLE, 0.9 + c * cell, 1.6 + r * cell, cell, cell,
                           colours[key], line=WHITE)
            sq.line.width = Pt(1)
    add_text(s, 0.9, 5.05, 4.4, 0.5, "Raster", size=26, bold=True, color=GREEN)
    add_text(s, 0.9, 5.6, 4.6, 1.2, "A grid of tiny coloured squares called pixels. "
             "A satellite photo is raster.", size=18)
    # Vector: shapes
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.1, 1.6, 4.4, 3.3, GREY_TINT).adjustments[0] = 0.05
    add_polyline(s, [(7.6, 2.1), (9.6, 1.95), (10.9, 2.9), (9.9, 4.2), (7.9, 3.9)], AMBER,
                 width=4, closed=True, fill=AMBER_TINT)
    add_polyline(s, [(7.3, 4.6), (8.6, 3.3), (10.2, 3.6), (11.3, 2.2)], BLUE, width=5)
    add_dot(s, 9.1, 2.8, 0.35, RED)
    add_text(s, 7.1, 5.05, 4.4, 0.5, "Vector", size=26, bold=True, color=AMBER)
    add_text(s, 7.1, 5.6, 5.0, 1.2, "Points, lines and polygons. The India outline you "
             "load in Exercise 1 is vector.", size=18)
    add_notes(s, "Zoom right into a satellite photo and it turns into squares: that is "
                 "raster. Vector data stores shapes as corner points, so it stays sharp "
                 "however far you zoom. In Exercise 1 the satellite basemap is raster and the "
                 "GeoJSON outlines are vector.")


def slide_parts_of_map(prs):
    s = new_slide(prs)
    add_title(s, "What is a map?")
    add_text(s, 0.6, 1.4, 6.2, 1.4,
             "A map is a picture of a place, drawn from above, that shows where things are "
             "and how they relate to each other.", size=20)
    # Map image: 1240 x 1754 pixels, shown 4.2 inches wide
    mx, my, mw = 8.45, 1.1, 4.2
    k = mw / 1240
    add_picture(s, HERE / "screenshot" / "Map-of-India.png", mx, my, w=mw)
    parts = [("Title", "says what the map shows", (390, 70)),
             ("North arrow", "shows which way is north", (1020, 220)),
             ("Outline", "a polygon layer on top", (300, 900)),
             ("Scale bar", "shows real distance", (930, 1660))]
    ty = 3.1
    for num, (name, meaning, (px, py)) in enumerate(parts, start=1):
        number_dot(s, 0.85, ty + 0.27, str(num))
        add_text(s, 1.3, ty, 2.2, 0.5, name, size=20, bold=True, color=GREEN)
        add_text(s, 3.5, ty, 4.4, 0.5, meaning, size=20)
        number_dot(s, mx + px * k, my + py * k, str(num))
        ty += 0.9
    add_notes(s, "Every good map answers three questions: what is it a map of (title), which "
                 "way is north (north arrow), and how big are things really (scale bar). Many "
                 "maps also have a legend that explains the colours and symbols. You will add "
                 "the title, north arrow and scale bar yourself in step 6 of Exercise 1.")


def slide_maps_tell_stories(prs):
    s = new_slide(prs)
    add_title(s, "Maps tell stories")
    add_picture(s, IMAGES / "lalkot_map.png", 0.6, 1.4, h=5.3)
    add_text(s, 6.9, 1.5, 5.9, 1.2,
             "This map shows the old walls of Lal Kot, a fort in Delhi, drawn on top of a "
             "satellite photo.", size=20)
    add_text(s, 6.9, 2.9, 5.9, 2.2,
             ["Yellow dots are points: monuments", "Dashed lines trace the old walls",
              "Coloured shapes are polygons: the palace area"], size=20, bullets=True)
    add_card(s, 6.9, 4.85, 5.8, 1.8, "Legend",
             "The box in the corner that explains what every colour and symbol means.",
             AMBER_TINT, RGBColor(0x92, 0x40, 0x0E))
    add_notes(s, "Archaeologists use satellite images and GIS to find and map old forts. "
                 "Point out the legend, the north arrow at top left and the scale bar. "
                 "Ask: what story is this map telling? (Where the old fort was, and what "
                 "is left of it today.)")


def slide_quiz(prs):
    s = new_slide(prs, dark=True)
    add_text(s, 0.6, 0.4, 12.1, 0.9, "Quick quiz", size=36, bold=True, color=WHITE)
    questions = ["What does GIS stand for?",
                 "Is the Moon a natural or an artificial satellite?",
                 "Give one example each of a point, a line and a polygon.",
                 "What does a scale bar tell you?"]
    y = 1.6
    for i, q in enumerate(questions, start=1):
        add_dot(s, 0.95, y + 0.28, 0.5, AMBER, outline=AMBER)
        add_text(s, 0.7, y + 0.02, 0.5, 0.5, str(i), size=20, bold=True, color=NAVY,
                 align=PP_ALIGN.CENTER)
        add_text(s, 1.45, y, 7.0, 0.8, q, size=22, color=WHITE)
        y += 1.0
    card = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.8, 1.6, 3.9, 3.9, GREEN)
    card.adjustments[0] = 0.08
    add_text(s, 9.1, 1.85, 3.3, 0.5, "NEXT UP", size=14, bold=True, color=AMBER_TINT)
    add_text(s, 9.1, 2.35, 3.3, 1.6, "Exercise 1: Map of India", size=28, bold=True,
             color=WHITE)
    add_text(s, 9.1, 3.85, 3.3, 1.5, "Open GeoLibre and build your own map, layer by layer.",
             size=18, color=WHITE)
    add_notes(s, "Answers: 1. Geographic Information System. 2. Natural. 3. Any sensible "
                 "examples, such as a school (point), a road (line), a lake (polygon). "
                 "4. How far apart things really are on the ground. Then hand out the "
                 "Exercise 1 tutorial.")


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for build in [slide_title, slide_journey, slide_remote_sensing, slide_how_satellites_see,
                  slide_satellites, slide_india_in_space, slide_satellite_photo, slide_change,
                  slide_land_cover, slide_bengaluru_story,
                  slide_gis, slide_points_lines_polygons, slide_spot_the_shapes,
                  slide_raster_vector, slide_parts_of_map, slide_maps_tell_stories,
                  slide_quiz]:
        build(prs)
    prs.save(OUTPUT)
    print(f"Saved {OUTPUT} with {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
